# rep_report.py
# WDI Visit Analytics Engine
# تقرير المندوب التفصيلي — ملف Excel من 9 شيتات برسوم Excel أصلية.

import io

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

from export_manager import (
    PRIMARY, SECONDARY, ACCENT, HEADER_FG,
    _header_fill, _header_font, _normal_font, _alt_fill, _border,
    _center, _right_align, _auto_width, _is_arabic_col,
)

RED   = "C00000"
AMBER = "FFC000"


def _block(ws, title: str, df: pd.DataFrame, row: int, color: str = PRIMARY,
           note: str = "") -> int:
    """اكتب عنواناً ثم جدولاً، وأعد رقم أول صف فارغ بعده."""
    ncols = max(1, len(df.columns) if df is not None and not df.empty else 1)
    if title:
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = Font(bold=True, color=HEADER_FG, name="Calibri", size=12)
        cell.fill = _header_fill(color)
        cell.alignment = _right_align()
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        ws.row_dimensions[row].height = 22
        row += 1
    if note:
        c = ws.cell(row=row, column=1, value=note)
        c.font = Font(italic=True, size=9, color="666666")
        row += 1
    if df is None or df.empty:
        ws.cell(row=row, column=1, value="لا توجد بيانات").font = _normal_font()
        return row + 2

    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=row, column=j, value=str(col))
        c.fill, c.font, c.border, c.alignment = _header_fill(color), _header_font(), _border(), _center()
    for i, rec in enumerate(df.itertuples(index=False), start=1):
        fill = _alt_fill() if i % 2 == 0 else PatternFill()
        for j, val in enumerate(rec, start=1):
            if val is None or (isinstance(val, float) and pd.isna(val)):
                val = ""
            c = ws.cell(row=row + i, column=j, value=val)
            c.fill, c.border, c.font = fill, _border(), _normal_font()
            c.alignment = _right_align() if _is_arabic_col(df.columns[j - 1]) else _center()
    return row + len(df) + 2


def _chart(ws, kind: str, title: str, anchor: str, header_row: int, last_row: int,
           label_col: int = 1, value_col: int = 2, height: float = 8, width: float = 16):
    """رسم Excel أصلي — يبقى حياً وقابلاً للتعديل داخل الملف."""
    if last_row <= header_row:
        return
    ch = {"bar": BarChart, "line": LineChart, "pie": PieChart}[kind]()
    if kind == "bar":
        ch.type = "col"
    ch.title = title
    ch.height, ch.width = height, width
    ch.add_data(Reference(ws, min_col=value_col, min_row=header_row, max_row=last_row),
                titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=label_col, min_row=header_row + 1, max_row=last_row))
    if kind != "pie":
        ch.legend = None
    ws.add_chart(ch, anchor)


def _sheet(wb, name: str, first: bool = False):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name
    ws.sheet_view.rightToLeft = True
    return ws


def export_rep_report(d: dict) -> bytes:
    """تقرير Excel شامل لمندوب واحد عن الفترة المختارة."""
    h = d["header"]
    wb = openpyxl.Workbook()

    # ── 1) الملخص ──
    ws = _sheet(wb, "الملخص", first=True)
    ws.merge_cells("A1:E1")
    t = ws.cell(row=1, column=1, value=f"تقرير أداء المندوب — {h['rep']}")
    t.font = Font(bold=True, color=HEADER_FG, name="Calibri", size=15)
    t.fill = _header_fill(PRIMARY)
    t.alignment = _center()
    ws.row_dimensions[1].height = 30
    for i, (k, v) in enumerate([("الفترة", h["period"]),
                                ("المحافظات المسؤول عنها", h["governorates"]),
                                ("محافظات زيارات الفترة", h.get("governorates_period", "—")),
                                ("مقارنة مع", h["prev_label"]),
                                ("تاريخ الإصدار", h["generated"])], start=2):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True, size=10)
        ws.cell(row=i, column=2, value=v).font = _normal_font()

    r = _block(ws, "المؤشرات مقارنةً بالفترة السابقة وبمتوسط الفريق", d["kpis"], 8)
    r = _block(ws, "نسب الأداء", d["ratios"], r, SECONDARY)
    geo_start = r
    r = _block(ws, "التوزيع الجغرافي لزيارات الفترة", d["geo_distribution"], r, ACCENT)
    if not d["geo_distribution"].empty:
        _chart(ws, "bar", "الزيارات حسب المحافظة", f"G{geo_start}",
               geo_start + 1, geo_start + len(d["geo_distribution"]) + 1)
    mix_start = r
    r = _block(ws, "توزيع حالات الزيارات في الفترة", d["status_mix"], r, SECONDARY)
    if not d["status_mix"].empty:
        _chart(ws, "pie", "توزيع حالات الزيارات", f"G{mix_start + 17}",
               mix_start + 1, mix_start + len(d["status_mix"]) + 1)
    _auto_width(ws)

    # ── 2) الزيارات ──
    ws = _sheet(wb, "الزيارات")
    r = _block(ws, "الزيارات شهرياً", d["visits_monthly"], 1)
    wk_start = r
    r = _block(ws, "الزيارات أسبوعياً", d["visits_weekly"], r, SECONDARY)
    if not d["visits_weekly"].empty:
        _chart(ws, "line", "اتجاه الزيارات أسبوعياً", f"G{wk_start}",
               wk_start + 1, wk_start + len(d["visits_weekly"]) + 1)
    r = _block(ws, "الزيارات يومياً", d["visits_daily"], r, ACCENT)
    wd_start = r
    r = _block(ws, "الزيارات حسب يوم الأسبوع", d["visits_weekday"], r, PRIMARY)
    if not d["visits_weekday"].empty:
        _chart(ws, "bar", "الزيارات حسب يوم الأسبوع", f"G{wd_start}",
               wd_start + 1, wd_start + len(d["visits_weekday"]) + 1)
    _auto_width(ws)

    # ── 3) غير المهتمين والمتوقفين ──
    ws = _sheet(wb, "غير مهتم ومتوقف")
    r = _block(ws, "ملخص زيارات العملاء غير المهتمين والمتوقفين", d["wasted_summary"], 1, RED,
               note="معدل التكرار = الزيارات ÷ العملاء · النسبة من إجمالي زيارات الفترة")
    if not d["wasted_summary"].empty:
        _chart(ws, "bar", "الزيارات غير المنتجة حسب الحالة", "H2", 3,
               2 + len(d["wasted_summary"]) + 1)
    r = _block(ws, "تفصيل كل زيارة", d["wasted_detail"], r, RED)
    _auto_width(ws)

    # ── 4) التحويلات ──
    ws = _sheet(wb, "التحويلات")
    r = _block(ws, "ملخص التحويلات إلى عميل حالي", d["conv_summary"], 1, ACCENT)
    r = _block(ws, "تفصيل تحويلات الفترة", d["conv_detail_period"], r, ACCENT)
    r = _block(ws, "كل تحويلات المندوب (تراكمي)", d["conv_detail_all"], r, SECONDARY)
    _auto_width(ws)

    # ── 5) أكثر العملاء زيارة ──
    ws = _sheet(wb, "أكثر العملاء زيارة")
    r = _block(ws, "أكثر 30 عميلاً زيارةً في الفترة — مع الحالة الحالية", d["top_customers"], 1)
    if not d["top_customers"].empty:
        _chart(ws, "bar", "أكثر العملاء زيارة", "J2", 2,
               2 + min(15, len(d["top_customers"])), height=10)
    _auto_width(ws)

    # ── 6) الوعود ──
    ws = _sheet(wb, "الوعود")
    r = _block(ws, "ملخص الوعود", d["promises_summary"], 1, AMBER,
               note="كل وعود المندوب عبر كل الفترات — غير مقيدة بالتاريخ المختار")
    r = _block(ws, "تفصيل الوعود", d["promises"], r, AMBER)
    _auto_width(ws)

    # ── 7) المنافسون ──
    ws = _sheet(wb, "المنافسون")
    r = _block(ws, "المنافسون في محافظات المندوب", d["competitors"], 1, RED,
               note="الأرقام تغطي كل تاريخ المندوب · عمود (ذُكر في الفترة) يخص الفترة المختارة")
    if not d["competitors"].empty:
        _chart(ws, "bar", "عدد العملاء لكل منافس", "H2", 3,
               2 + min(12, len(d["competitors"])) + 1)
    r = _block(ws, "مصفوفة المنافسين حسب المحافظة", d["competitor_matrix"], r, RED)
    r = _block(ws, "عملاء مهددون مرتبطون بمنافس", d["competitor_losing"], r, RED)
    _auto_width(ws)

    # ── 8) عملاء مهملون ──
    ws = _sheet(wb, "عملاء مهملون")
    r = _block(ws, "عملاء من محفظة المندوب بلا زيارة", d["neglected_summary"], 1, AMBER,
               note="محسوبة من تاريخ آخر زيارة لأي مندوب — أولوية المتابعة")
    if not d["neglected_summary"].empty:
        _chart(ws, "bar", "العملاء المهملون حسب مدة الانقطاع", "F2", 3,
               2 + len(d["neglected_summary"]) + 1)
    r = _block(ws, "التفصيل (الأطول انقطاعاً أولاً)", d["neglected_detail"], r, AMBER)
    _auto_width(ws)

    # ── 9) جودة الأداء ──
    ws = _sheet(wb, "جودة الأداء")
    r = _block(ws, "حركة محفظة العملاء خلال الفترة", d["portfolio_moves"], 1, SECONDARY,
               note="تغيّر حالة العملاء على يد هذا المندوب")
    r = _block(ws, "جودة تسجيل الملاحظات", d["note_quality"], r, SECONDARY)
    r = _block(ws, "زيارات لم تتم (غير موجود / مسافر)", d["nomeeting_detail"], r, RED)
    r = _block(ws, "زيارات مكررة لنفس العميل في نفس اليوم", d["same_day_dupes"], r, RED)
    _auto_width(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
