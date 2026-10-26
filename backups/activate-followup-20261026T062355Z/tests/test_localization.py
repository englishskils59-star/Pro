import copy
import io
import unittest
from unittest.mock import patch
import pandas as pd
import plotly.graph_objects as go
from openpyxl import load_workbook
import localization as loc
from export_manager import export_classification_results, export_monthly_report
from rep_report import _block, _sheet
import openpyxl

class LocalizationTests(unittest.TestCase):
    def setUp(self):
        self.state={}
        self.patch=patch.object(loc._st,'session_state',self.state)
        self.patch.start()
    def tearDown(self): self.patch.stop()
    def test_arabic_default_and_unknown_fallback(self):
        self.assertEqual(loc.language(),'ar')
        self.assertEqual(loc.tr('تصنيف العملاء'),'تصنيف العملاء')
        self.state['_ui_language']='en'
        self.assertEqual(loc.tr('تصنيف العملاء'),'Customer Classification')
        self.assertEqual(loc.tr('المندوب عبارة جديدة غير مترجمة'),'المندوب عبارة جديدة غير مترجمة')
    def test_inserted_values_are_not_translated(self):
        self.state['_ui_language']='en'
        path='D:/المندوب/العملاء'
        self.assertEqual(loc.translate_message('✅ تم تعيين مسار البيانات: '+path),'✅ Data Path Set: '+path)
        text=loc.translated('العميل: المندوب')
        self.assertEqual(loc.tr(text),text)
    def test_compound_options_keep_business_names_and_canonical_values(self):
        self.state['_ui_language']='en'
        raw='العميل — المنيا (2 اسم)'
        self.assertEqual(loc.option_label('merge_grp',raw),'العميل — المنيا (2 names)')
        self.assertEqual(raw,'العميل — المنيا (2 اسم)')
        self.assertEqual(loc.option_label('del_rules','المندوب → عميل حالي (60)'),
                         'المندوب → Current Customer (60)')
    def test_display_copy_preserves_records_and_schema(self):
        self.state['_ui_language']='en'
        df=pd.DataFrame({'العميل':['المندوب'],'Visit Notes':['محتمل من القاهرة'],'المؤشر':['إجمالي الزيارات']})
        before=df.copy(deep=True)
        shown=loc.display_frame(df)
        pd.testing.assert_frame_equal(df,before)
        self.assertEqual(shown['Customer'][0],'المندوب')
        self.assertEqual(shown['Visit Notes'][0],'محتمل من القاهرة')
        self.assertEqual(shown['Metric'][0],'Total Visits')
    def test_chart_copy_does_not_mutate_cached_figure(self):
        fig=go.Figure(go.Pie(labels=['عميل حالي','محتمل'],values=[2,3]))
        before=fig.to_json()
        self.state['_ui_language']='en'
        shown=loc.display_figure(fig)
        self.assertEqual(list(shown.data[0].labels),['Current Customer','Potential'])
        self.assertEqual(fig.to_json(),before)
    def test_chart_business_names_and_values_remain_intact(self):
        self.state['_ui_language']='en'
        fig=go.Figure(go.Bar(x=[12],y=['محتمل'],orientation='h',hovertemplate='Sales Rep Name=%{y}'))
        before=fig.to_json()
        shown=loc.display_figure(fig)
        self.assertEqual(list(shown.data[0].y),['محتمل'])
        self.assertEqual(list(shown.data[0].x),[12])
        self.assertEqual(fig.to_json(),before)
    def test_raw_export_schema_and_records_do_not_change(self):
        df=pd.DataFrame({'Customer Name':['عميل حالي'],'Visit Notes':['المندوب كتب ملاحظة'],'Display Status':['Current Customer']})
        books=[]
        for lang in ['ar','en']:
            self.state['_ui_language']=lang
            books.append(load_workbook(io.BytesIO(export_classification_results(df))))
        rows=[list(w.active.values) for w in books]
        self.assertEqual(rows[0][1:],rows[1][1:])
    def test_human_reports_follow_language(self):
        self.state['_ui_language']='en'
        wb=openpyxl.Workbook();ws=_sheet(wb,'الملخص',first=True)
        _block(ws,'الوعود',pd.DataFrame({'العميل':['عميل حالي'],'المؤشر':['إجمالي الزيارات']}),1)
        self.assertEqual(ws.title,'Summary')
        self.assertFalse(ws.sheet_view.rightToLeft)
        self.assertEqual(ws.cell(1,1).value,'Promises')
        self.assertEqual(ws.cell(3,1).value,'عميل حالي')
        self.assertEqual(ws.cell(3,2).value,'Total Visits')
        report=load_workbook(io.BytesIO(export_monthly_report({'metrics':[('إجمالي الزيارات',10,5)]},{})))
        self.assertEqual(report.active.title,'Monthly Comparison')
        self.assertFalse(report.active.sheet_view.rightToLeft)

if __name__=='__main__': unittest.main()
